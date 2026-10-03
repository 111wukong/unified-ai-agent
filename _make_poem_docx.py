# -*- coding: utf-8 -*-
"""在桌面生成一首古诗的 Word 文档，排版工整。"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

TITLE = "秋夜书怀"
AUTHOR = "（当代）佚名"
LINES = [
    "霜天寥落雁南飞，",
    "独倚危楼看落晖。",
    "一夜西风吹不尽，",
    "满城黄叶送秋归。",
]

OUT = "/Users/wukong/Desktop/古诗·秋夜书怀.docx"


def set_font(run, name, size, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


doc = Document()

# 页面正文默认字体
style = doc.styles["Normal"]
style.font.name = "楷体"
style.font.size = Pt(16)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")

# 标题
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
set_font(p.add_run(TITLE), "黑体", 26, bold=True)

# 署名
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(18)
set_font(p.add_run(AUTHOR), "楷体", 12)

# 正文四句，逐句居中，行距统一
for line in LINES:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = Pt(34)
    pf.space_after = Pt(0)
    set_font(p.add_run(line), "楷体", 18)

doc.save(OUT)
print("saved:", OUT)
