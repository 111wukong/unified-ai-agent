# -*- coding: utf-8 -*-
"""校验桌面上的古诗 Word 文档内容与排版。"""
import docx

path = "/Users/wukong/Desktop/古诗·秋夜书怀.docx"
d = docx.Document(path)
print("段落数:", len(d.paragraphs))
for i, p in enumerate(d.paragraphs):
    print(i, repr(p.text), "align=", p.alignment)
